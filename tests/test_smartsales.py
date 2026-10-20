import csv
from decimal import Decimal
from datetime import datetime
from openpyxl import Workbook

from src.smartsales_stage3 import (
    read_sales_file,
    prepare_records,
    calculate_summary,
    find_duplicate_ids,
)

def test_basic_math():
    """Check that Python calculates a simple revenue amount correctly."""
    quantity = 4
    unit_price = 5.00

    revenue = quantity * unit_price

    assert revenue == 20.00


def test_read_sales_file_reads_csv(tmp_path):
    """Check that a CSV file is read correctly."""
    csv_file = tmp_path / "sales.csv"

    with csv_file.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=[
                "order_id",
                "order_date",
                "product",
                "category",
                "quantity",
                "unit_price",
            ],
        )
        writer.writeheader()
        writer.writerow({
            "order_id": "1001",
            "order_date": "2026-10-01",
            "product": "Notebook",
            "category": "Accessories",
            "quantity": "2",
            "unit_price": "5.00",
        })

    rows = read_sales_file(csv_file)

    assert rows is not None
    assert len(rows) == 1
    assert rows[0][0] == 2
    assert rows[0][1]["product"] == "Notebook"
    assert rows[0][1]["quantity"] == "2"


def test_read_sales_file_rejects_missing_columns(tmp_path):
    """Check that a CSV missing required columns is rejected."""
    csv_file = tmp_path / "incomplete.csv"
    csv_file.write_text(
        "order_id,product\n1001,Notebook\n",
        encoding="utf-8",
    )

    rows = read_sales_file(csv_file)

    assert rows is None


def test_prepare_records_separates_valid_and_invalid_rows(tmp_path):
    """Check that valid sales are accepted and invalid sales are rejected."""
    csv_file = tmp_path / "validation.csv"

    with csv_file.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=[
                "order_id",
                "order_date",
                "product",
                "category",
                "quantity",
                "unit_price",
            ],
        )
        writer.writeheader()

        # Valid sales record.
        writer.writerow({
            "order_id": "1001",
            "order_date": "2026-10-01",
            "product": "Notebook",
            "category": "Accessories",
            "quantity": "2",
            "unit_price": "5.00",
        })

        # Invalid record: quantity cannot be negative.
        writer.writerow({
            "order_id": "1002",
            "order_date": "2026-10-02",
            "product": "Pen",
            "category": "Accessories",
            "quantity": "-3",
            "unit_price": "1.50",
        })

    rows = read_sales_file(csv_file)
    assert rows is not None

    valid_records, invalid_records = prepare_records(rows)

    assert len(valid_records) == 1
    assert len(invalid_records) == 1

    assert valid_records[0]["product"] == "Notebook"
    assert valid_records[0]["quantity"] == 2
    assert valid_records[0]["unit_price"] == Decimal("5.00")

    assert "Quantity must be a positive integer" in invalid_records[0]["errors"]


def test_calculate_summary_returns_correct_totals():
    """Check total revenue, units sold, and revenue by product."""
    valid_records = [
        {
            "product": "Notebook",
            "quantity": 2,
            "unit_price": Decimal("5.00"),
        },
        {
            "product": "Pen",
            "quantity": 3,
            "unit_price": Decimal("1.50"),
        },
        {
            "product": "Backpack",
            "quantity": 1,
            "unit_price": Decimal("25.00"),
        },
    ]

    total_revenue, total_units, revenue_by_product = calculate_summary(
        valid_records
    )

    assert total_revenue == Decimal("39.50")
    assert total_units == 6
    assert revenue_by_product["Notebook"] == Decimal("10.00")
    assert revenue_by_product["Pen"] == Decimal("4.50")
    assert revenue_by_product["Backpack"] == Decimal("25.00")


def test_calculate_summary_handles_empty_records():
    """Check that an empty sales list produces zero totals."""
    total_revenue, total_units, revenue_by_product = calculate_summary([])

    assert total_revenue == Decimal("0.00")
    assert total_units == 0
    assert revenue_by_product == {}


def test_calculate_summary_combines_revenue_for_same_product():
    """Check that multiple sales of one product are added together."""
    valid_records = [
        {
            "product": "Notebook",
            "quantity": 2,
            "unit_price": Decimal("5.00"),
        },
        {
            "product": "Notebook",
            "quantity": 4,
            "unit_price": Decimal("5.00"),
        },
    ]

    total_revenue, total_units, revenue_by_product = calculate_summary(
        valid_records
    )

    assert total_revenue == Decimal("30.00")
    assert total_units == 6
    assert revenue_by_product["Notebook"] == Decimal("30.00")


def test_invalid_records_are_excluded_from_revenue(tmp_path):
    """Check that invalid sales are excluded from financial totals."""
    import csv

    csv_file = tmp_path / "mixed_sales.csv"

    with csv_file.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=[
                "order_id",
                "order_date",
                "product",
                "category",
                "quantity",
                "unit_price",
            ],
        )
        writer.writeheader()

        writer.writerow({
            "order_id": "1001",
            "order_date": "2026-10-01",
            "product": "Notebook",
            "category": "Accessories",
            "quantity": "2",
            "unit_price": "5.00",
        })

        writer.writerow({
            "order_id": "1002",
            "order_date": "2026-10-02",
            "product": "Pen",
            "category": "Accessories",
            "quantity": "-3",
            "unit_price": "100.00",
        })

    rows = read_sales_file(csv_file)
    assert rows is not None

    valid_records, invalid_records = prepare_records(rows)

    total_revenue, total_units, revenue_by_product = calculate_summary(
        valid_records
    )

    assert len(valid_records) == 1
    assert len(invalid_records) == 1
    assert total_revenue == Decimal("10.00")
    assert total_units == 2
    assert revenue_by_product == {"Notebook": Decimal("10.00")}


def test_invalid_date_is_rejected(tmp_path):
    """A row with an invalid date should be rejected."""
    csv_file = tmp_path / "invalid_date.csv"
    csv_file.write_text(
        "order_id,order_date,product,category,quantity,unit_price\n"
        "1001,not-a-date,Notebook,Accessories,2,5.00\n",
        encoding="utf-8",
    )

    rows = read_sales_file(csv_file)
    assert rows is not None

    valid_records, invalid_records = prepare_records(rows)

    assert len(valid_records) == 0
    assert len(invalid_records) == 1
    assert "date" in invalid_records[0]["errors"].lower()


def test_zero_quantity_is_rejected(tmp_path):
    """Quantity must be a positive integer."""
    csv_file = tmp_path / "zero_quantity.csv"
    csv_file.write_text(
        "order_id,order_date,product,category,quantity,unit_price\n"
        "1001,2026-10-01,Notebook,Accessories,0,5.00\n",
        encoding="utf-8",
    )

    rows = read_sales_file(csv_file)
    assert rows is not None

    valid_records, invalid_records = prepare_records(rows)

    assert len(valid_records) == 0
    assert len(invalid_records) == 1
    assert "quantity" in invalid_records[0]["errors"].lower()


def test_invalid_unit_price_is_rejected(tmp_path):
    """A non-numeric unit price should be rejected."""
    csv_file = tmp_path / "invalid_price.csv"
    csv_file.write_text(
        "order_id,order_date,product,category,quantity,unit_price\n"
        "1001,2026-10-01,Notebook,Accessories,2,abc\n",
        encoding="utf-8",
    )

    rows = read_sales_file(csv_file)
    assert rows is not None

    valid_records, invalid_records = prepare_records(rows)

    assert len(valid_records) == 0
    assert len(invalid_records) == 1
    assert "price" in invalid_records[0]["errors"].lower()


def test_duplicate_order_ids_are_rejected(tmp_path):
    """Every row sharing a duplicate order ID should be rejected."""
    csv_file = tmp_path / "duplicate_ids.csv"
    csv_file.write_text(
        "order_id,order_date,product,category,quantity,unit_price\n"
        "1001,2026-10-01,Notebook,Accessories,2,5.00\n"
        "1001,2026-10-02,Pen,Accessories,3,1.50\n",
        encoding="utf-8",
    )

    rows = read_sales_file(csv_file)
    assert rows is not None

    duplicate_ids = find_duplicate_ids(rows)
    valid_records, invalid_records = prepare_records(rows)

    assert "1001" in duplicate_ids
    assert len(valid_records) == 0
    assert len(invalid_records) == 2


def test_invalid_category_is_rejected(tmp_path):
    """An unsupported category should be rejected."""
    csv_file = tmp_path / "invalid_category.csv"
    csv_file.write_text(
        "order_id,order_date,product,category,quantity,unit_price\n"
        "1001,2026-10-01,Notebook,UnknownCategory,2,5.00\n",
        encoding="utf-8",
    )

    rows = read_sales_file(csv_file)
    assert rows is not None

    valid_records, invalid_records = prepare_records(rows)

    assert len(valid_records) == 0
    assert len(invalid_records) == 1
    assert "category" in invalid_records[0]["errors"].lower()


from openpyxl import load_workbook

from src.smartsales_stage3 import create_excel_report


def test_excel_report_is_created(tmp_path):
    """The report should be saved as an Excel workbook."""
    output_file = tmp_path / "reports" / "test_report.xlsx"

    rows = [
        (2, {
            "order_id": "1001",
            "order_date": "2026-10-01",
            "product": "Notebook",
            "category": "Accessories",
            "quantity": "2",
            "unit_price": "5.00",
        })
    ]

    valid_records = [
        {
            "order_id": "1001",
            "order_date": "2026-10-01",
            "product": "Notebook",
            "category": "Accessories",
            "quantity": 2,
            "unit_price": Decimal("5.00"),
        }
    ]

    invalid_records = []

    report_path, total_revenue, total_units = create_excel_report(
        rows,
        valid_records,
        invalid_records,
        output_file,
    )

    assert report_path.exists()
    assert report_path.suffix == ".xlsx"
    assert total_revenue == Decimal("10.00")
    assert total_units == 2


def test_excel_report_has_expected_worksheets(tmp_path):
    """The workbook should contain all three report worksheets."""
    output_file = tmp_path / "test_report.xlsx"

    rows = []
    valid_records = []
    invalid_records = []

    create_excel_report(
        rows,
        valid_records,
        invalid_records,
        output_file,
    )

    workbook = load_workbook(output_file, read_only=True)

    assert workbook.sheetnames == [
        "Sales Summary",
        "Product Revenue",
        "Invalid Records",
    ]

    workbook.close()


def test_excel_summary_contains_correct_totals(tmp_path):
    """The summary worksheet should show correct revenue and unit totals."""
    output_file = tmp_path / "summary_report.xlsx"

    rows = [
        (2, {
            "order_id": "1001",
            "order_date": "2026-10-01",
            "product": "Notebook",
            "category": "Accessories",
            "quantity": "2",
            "unit_price": "5.00",
        }),
        (3, {
            "order_id": "1002",
            "order_date": "2026-10-02",
            "product": "Pen",
            "category": "Accessories",
            "quantity": "3",
            "unit_price": "1.50",
        }),
    ]

    valid_records = [
        {
            "order_id": "1001",
            "order_date": "2026-10-01",
            "product": "Notebook",
            "category": "Accessories",
            "quantity": 2,
            "unit_price": Decimal("5.00"),
        },
        {
            "order_id": "1002",
            "order_date": "2026-10-02",
            "product": "Pen",
            "category": "Accessories",
            "quantity": 3,
            "unit_price": Decimal("1.50"),
        },
    ]

    create_excel_report(rows, valid_records, [], output_file)

    workbook = load_workbook(output_file, data_only=True)
    summary = workbook["Sales Summary"]

    assert summary["B2"].value == 2
    assert summary["B3"].value == 2
    assert summary["B4"].value == 0
    assert summary["B5"].value == 5
    assert summary["B6"].value == 14.5

    workbook.close()


def test_invalid_records_appear_in_excel_report(tmp_path):
    """Invalid rows should be listed but excluded from sales totals."""
    output_file = tmp_path / "invalid_records_report.xlsx"

    rows = [
        (2, {
            "order_id": "1001",
            "order_date": "2026-10-01",
            "product": "Notebook",
            "category": "Accessories",
            "quantity": "2",
            "unit_price": "5.00",
        }),
        (3, {
            "order_id": "1002",
            "order_date": "2026-10-02",
            "product": "Pen",
            "category": "Accessories",
            "quantity": "-3",
            "unit_price": "100.00",
        }),
    ]

    valid_records = [
        {
            "order_id": "1001",
            "order_date": "2026-10-01",
            "product": "Notebook",
            "category": "Accessories",
            "quantity": 2,
            "unit_price": Decimal("5.00"),
        }
    ]

    invalid_records = [
        {
            "csv_row": 3,
            "order_id": "1002",
            "order_date": "2026-10-02",
            "product": "Pen",
            "category": "Accessories",
            "quantity": "-3",
            "unit_price": "100.00",
            "errors": "Quantity must be a positive integer",
        }
    ]

    create_excel_report(
        rows,
        valid_records,
        invalid_records,
        output_file,
    )

    workbook = load_workbook(output_file, data_only=True)

    summary = workbook["Sales Summary"]
    errors_sheet = workbook["Invalid Records"]

    # Only the valid order contributes to revenue.
    assert summary["B6"].value == 10.0
    assert summary["B3"].value == 1
    assert summary["B4"].value == 1

    # The invalid order should still appear in the error worksheet.
    assert errors_sheet["B2"].value == "1002"
    assert errors_sheet["D2"].value == "Pen"
    assert "Quantity must be a positive integer" in errors_sheet["H2"].value

    workbook.close()


def test_empty_csv_file_is_rejected(tmp_path):
    """An empty CSV file should return None."""
    csv_file = tmp_path / "empty.csv"
    csv_file.write_text("", encoding="utf-8")

    rows = read_sales_file(csv_file)

    assert rows is None


def test_csv_with_header_only_returns_no_records(tmp_path):
    """A CSV with valid headers but no data should return an empty list."""
    csv_file = tmp_path / "header_only.csv"
    csv_file.write_text(
        "order_id,order_date,product,category,quantity,unit_price\n",
        encoding="utf-8",
    )

    rows = read_sales_file(csv_file)

    assert rows == []


def test_negative_unit_price_is_rejected(tmp_path):
    """A negative unit price must not be accepted."""
    csv_file = tmp_path / "negative_price.csv"
    csv_file.write_text(
        "order_id,order_date,product,category,quantity,unit_price\n"
        "1001,2026-10-01,Notebook,Accessories,2,-5.00\n",
        encoding="utf-8",
    )

    rows = read_sales_file(csv_file)
    assert rows is not None

    valid_records, invalid_records = prepare_records(rows)

    assert len(valid_records) == 0
    assert len(invalid_records) == 1
    assert "price" in invalid_records[0]["errors"].lower()


def test_all_rows_with_duplicate_ids_are_invalid(tmp_path):
    """Neither row should be accepted when both have the same order ID."""
    csv_file = tmp_path / "duplicate_orders.csv"
    csv_file.write_text(
        "order_id,order_date,product,category,quantity,unit_price\n"
        "1001,2026-10-01,Notebook,Accessories,2,5.00\n"
        "1001,2026-10-02,Pen,Accessories,3,1.50\n",
        encoding="utf-8",
    )

    rows = read_sales_file(csv_file)
    assert rows is not None

    valid_records, invalid_records = prepare_records(rows)

    assert len(valid_records) == 0
    assert len(invalid_records) == 2

    for record in invalid_records:
        assert "duplicate" in record["errors"].lower()


def test_large_revenue_is_calculated_accurately():
    """Large monetary totals should retain exact decimal arithmetic."""
    records = [
        {
            "order_id": "1001",
            "order_date": "2026-10-01",
            "product": "Industrial Equipment",
            "category": "Electronics",
            "quantity": 1000000,
            "unit_price": Decimal("999999.99"),
        }
    ]

    total_revenue, total_units, revenue_by_product = calculate_summary(records)

    assert total_revenue == Decimal("999999990000.00")
    assert total_units == 1000000
    assert revenue_by_product["Industrial Equipment"] == Decimal("999999990000.00")


def test_csv_row_with_missing_value_is_handled(tmp_path):
    """A missing cell should not crash CSV loading."""
    csv_file = tmp_path / "missing_value.csv"

    csv_file.write_text(
        "order_id,order_date,product,category,quantity,unit_price\n"
        "1001,2026-10-01,Notebook,Accessories,2\n",
        encoding="utf-8",
    )

    rows = read_sales_file(csv_file)

    assert rows is not None
    assert len(rows) == 1
    assert rows[0][1]["unit_price"] == ""

    valid_records, invalid_records = prepare_records(rows)

    assert len(valid_records) == 0
    assert len(invalid_records) == 1
    assert "price" in invalid_records[0]["errors"].lower()



def test_csv_row_with_extra_value_is_rejected(
    tmp_path, capsys
):
    """A row with extra values should be rejected clearly."""
    csv_file = tmp_path / "extra_value.csv"

    csv_file.write_text(
        "order_id,order_date,product,category,quantity,unit_price\n"
        "1001,2026-10-01,Notebook,Accessories,2,5.00,EXTRA\n",
        encoding="utf-8",
    )

    rows = read_sales_file(csv_file)
    captured = capsys.readouterr()

    assert rows is None
    assert "row 2" in captured.out.lower()
    assert "more values than expected" in captured.out.lower()


def test_non_finite_unit_prices_are_rejected(tmp_path):
    """NaN and infinite prices must never enter revenue totals."""
    csv_file = tmp_path / "non_finite_prices.csv"

    csv_file.write_text(
        "order_id,order_date,product,category,quantity,unit_price\n"
        "1001,2026-10-01,Notebook,Accessories,2,NaN\n"
        "1002,2026-10-01,Pen,Accessories,1,Infinity\n"
        "1003,2026-10-01,Folder,Accessories,3,-Infinity\n",
        encoding="utf-8",
    )

    rows = read_sales_file(csv_file)

    assert rows is not None

    valid_records, invalid_records = prepare_records(rows)

    assert len(valid_records) == 0
    assert len(invalid_records) == 3

    total_revenue, total_units, revenue_by_product = (
        calculate_summary(valid_records)
    )

    assert total_revenue == Decimal("0.00")
    assert total_units == 0
    assert revenue_by_product == {}


def test_read_sales_file_reads_excel(tmp_path):
    """Check that an Excel workbook is read correctly."""
    excel_file = tmp_path / "sales.xlsx"
    workbook = Workbook()
    worksheet = workbook.active

    worksheet.append([
        "order_id",
        "order_date",
        "product",
        "category",
        "quantity",
        "unit_price",
    ])
    worksheet.append([
        1001,
        datetime(2026, 10, 1),
        "Notebook",
        "Accessories",
        2,
        5.00,
    ])

    workbook.save(excel_file)
    workbook.close()

    rows = read_sales_file(excel_file)

    assert rows is not None
    assert len(rows) == 1
    assert rows[0][0] == 2
    assert rows[0][1]["order_id"] == "1001"
    assert rows[0][1]["order_date"] == "2026-10-01"
    assert rows[0][1]["product"] == "Notebook"
    assert rows[0][1]["quantity"] == "2"


def test_read_sales_file_rejects_excel_missing_columns(tmp_path, capsys):
    """Check that an Excel file with missing columns is rejected."""
    excel_file = tmp_path / "incomplete.xlsx"
    workbook = Workbook()
    worksheet = workbook.active

    worksheet.append(["order_id", "product"])
    worksheet.append([1001, "Notebook"])

    workbook.save(excel_file)
    workbook.close()

    rows = read_sales_file(excel_file)

    captured = capsys.readouterr()

    assert rows is None
    assert "Missing required columns" in captured.out


def test_read_sales_file_rejects_unsupported_extension(tmp_path, capsys):
    """Check that unsupported file extensions are rejected."""
    unsupported_file = tmp_path / "sales.xls"
    unsupported_file.write_text("sample", encoding="utf-8")

    rows = read_sales_file(unsupported_file)

    captured = capsys.readouterr()

    assert rows is None
    assert "Unsupported file type" in captured.out


def test_load_config_uses_defaults_when_file_is_missing(tmp_path):
    """Missing configuration files should use built-in defaults."""
    from src.smartsales_stage3 import (
        DEFAULT_INPUT,
        DEFAULT_OUTPUT,
        load_config,
    )

    settings = load_config(tmp_path / "missing.json")

    assert settings["input_file"] == str(DEFAULT_INPUT)
    assert settings["output_file"] == str(DEFAULT_OUTPUT)
    assert settings["report_title"] == "SmartSales Report"
    assert settings["currency_symbol"] == "$"
    assert settings["top_products_count"] == 5


def test_load_config_reads_custom_settings(tmp_path):
    """Custom paths should be loaded from JSON."""
    import json
    from src.smartsales_stage3 import load_config

    config_file = tmp_path / "config.json"
    config_file.write_text(
        json.dumps({
            "input_file": "data/custom_sales.csv",
            "output_file": "reports/custom_report.xlsx",
            "report_title": "Monthly Sales",
            "currency_symbol": "€",
            "top_products_count": 3,
        }),
        encoding="utf-8",
    )

    settings = load_config(config_file)

    assert settings["input_file"] == "data/custom_sales.csv"
    assert settings["output_file"] == "reports/custom_report.xlsx"
    assert settings["report_title"] == "Monthly Sales"
    assert settings["currency_symbol"] == "€"
    assert settings["top_products_count"] == 3



def test_load_config_rejects_invalid_json(tmp_path):
    """Invalid JSON should produce a clear configuration error."""
    import pytest
    from src.smartsales_stage3 import load_config

    config_file = tmp_path / "config.json"
    config_file.write_text("{invalid json", encoding="utf-8")

    with pytest.raises(ValueError, match="Unable to load configuration"):
        load_config(config_file)


def test_load_config_rejects_empty_path(tmp_path):
    """Configuration paths must be non-empty strings."""
    import json
    import pytest
    from src.smartsales_stage3 import load_config

    config_file = tmp_path / "config.json"
    config_file.write_text(
        json.dumps({"input_file": "  "}),
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="input_file"):
        load_config(config_file)



def test_load_config_rejects_invalid_top_products_count(tmp_path):
    """The top-products count must be a positive integer."""
    import json
    import pytest
    from src.smartsales_stage3 import load_config

    config_file = tmp_path / "config.json"
    config_file.write_text(
        json.dumps({"top_products_count": 0}),
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="top_products_count"):
        load_config(config_file)

def test_summarize_calculates_average_order_value():
    """Multiple product lines for one order count as one order."""
    import pandas as pd
    from app import summarize

    sales = pd.DataFrame([
        {"order_id": 1, "product": "Notebook", "quantity": 2, "revenue": 10.0},
        {"order_id": 1, "product": "Pen", "quantity": 3, "revenue": 6.0},
        {"order_id": 2, "product": "Bag", "quantity": 1, "revenue": 24.0},
    ])

    result = summarize(sales)

    assert result["transactions"] == 3
    assert result["orders"] == 2
    assert result["average"] == 40.0 / 3
    assert result["average_order_value"] == 20.0


def test_business_insights_calculates_category_shares():
    """Category shares should add up to 100 percent."""
    import pandas as pd
    from app import calculate_business_insights

    sales = pd.DataFrame([
        {"product": "Notebook", "category": "Stationery", "revenue": 30.0},
        {"product": "Pen", "category": "Stationery", "revenue": 20.0},
        {"product": "Bag", "category": "Accessories", "revenue": 50.0},
    ])

    result = calculate_business_insights(sales)
    shares = result["category_shares"].set_index("category")["share_percent"]

    assert shares["Stationery"] == 50.0
    assert shares["Accessories"] == 50.0
    assert round(shares.sum(), 5) == 100.0


def test_business_insights_calculates_product_revenue_shares():
    """Top-product shares should use total revenue as the denominator."""
    import pandas as pd
    from app import calculate_business_insights

    sales = pd.DataFrame([
        {"product": "Notebook", "category": "Stationery", "revenue": 60.0},
        {"product": "Pen", "category": "Stationery", "revenue": 30.0},
        {"product": "Bag", "category": "Accessories", "revenue": 10.0},
    ])

    result = calculate_business_insights(sales, top_products_count=2)

    assert result["top_product_share"] == 60.0
    assert result["top_products_share"] == 90.0


def test_business_insights_handles_empty_data():
    """Empty sales data should return zero shares without errors."""
    import pandas as pd
    from app import calculate_business_insights

    result = calculate_business_insights(pd.DataFrame())

    assert result["category_shares"].empty
    assert result["top_product_share"] == 0.0
    assert result["top_products_share"] == 0.0